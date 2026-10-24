"""Creates a formatted .xlsx of all site text, to upload as the editable Google Sheet.

python3 build/make_sheet.py out.xlsx
Rows without an ID (section titles) are ignored by the site.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.path.insert(0, str(Path(__file__).parent))
import sheet  # noqa: E402

GROUPS = [
    ("Browser tab & menu", ["title", "description", "nav", "cta"]),
    ("Header (top of the page)", ["hello", "h1", "tagline", "see_work", "showcase_label", "showcase", "pillars"]),
    ("Brands row", ["brands_label", "settings.brands"]),
    ("What I bring", ["bring_eyebrow", "bring_h2", "bring_lead", "tiles"]),
    ("Projects", ["work_eyebrow", "work_h2", "work_lead", "cases"]),
    ("I also write", ["writing_label", "writing"]),
    ("Where I'm heading", ["dir_eyebrow", "dir_quote", "dir_points"]),
    ("Experience", ["exp_eyebrow", "exp_h2", "jobs", "earlier"]),
    ("About me", ["about_eyebrow", "about_h2", "about", "facts"]),
    ("How I work", ["proc_eyebrow", "proc_h2", "steps"]),
    ("Contact & footer", ["contact_eyebrow", "contact_h2", "contact_lead", "linkedin_btn", "email_btn",
                          "settings.email", "settings.linkedin", "footer"]),
]

MERLOT, CREAM, ICE, GREY = "561D24", "F5EFE6", "D6E3EF", "8A8A8A"


def short_label(label):
    # The section row already says where we are, so drop the "Section:" prefix
    head, sep, rest = label.partition(": ")
    if sep and not head.startswith(("Project", "Settings")):
        return rest[:1].upper() + rest[1:]
    return label


def build(out):
    rows = {key: (lab, vals) for key, lab, vals in sheet.rows()}
    used = set()
    wb = Workbook()
    ws = wb.active
    ws.title = "Website text"
    ws.append(["What it is", "English", "Swedish", "Finnish", "ID (don't edit)"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=MERLOT)
        c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26

    thin = Border(bottom=Side(style="thin", color="E4DCD0"))
    for n, (title, prefixes) in enumerate(GROUPS, 1):
        ws.append([f"{n}. {title.upper()}", "", "", "", ""])
        r = ws.max_row
        for c in ws[r]:
            c.fill = PatternFill("solid", fgColor=ICE)
            c.font = Font(bold=True, color=MERLOT, size=12)
            c.alignment = Alignment(vertical="center")
        ws.row_dimensions[r].height = 28
        for key, (lab, vals) in rows.items():
            top = key if key.startswith("settings.") else key.split(".")[0]
            if top not in prefixes or key in used:
                continue
            used.add(key)
            ws.append([short_label(lab), *vals, key])
            r = ws.max_row
            for i, c in enumerate(ws[r]):
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.border = thin
                if i == 0:
                    c.fill = PatternFill("solid", fgColor=CREAM)
                    c.font = Font(bold=True, color=MERLOT)
                if i == 4:
                    c.font = Font(color=GREY, size=9)
    missing = [k for k in rows if k not in used]
    assert not missing, f"rows not placed in a section: {missing}"

    for col, width in zip("ABCDE", [30, 55, 55, 55, 18]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "B2"

    help_ws = wb.create_sheet("How to use")
    lines = [
        "How to edit the website",
        "",
        "1. Go to the 'Website text' tab. Blue rows are the page sections, in the same order as on the website.",
        "2. Change the text in the English, Swedish or Finnish column.",
        "3. The website updates by itself within about 10–15 minutes.",
        "",
        "Good to know",
        "• Don't change the grey ID column. It tells the website where each text goes.",
        "• An empty cell keeps the current text on the website.",
        "• Text inside <em> … </em> in the headline is shown in the accent colour.",
        "• <br> makes a line break.",
        "• Adding new rows does nothing on the website. Ask Claude if you want new sections or items.",
    ]
    for line in lines:
        help_ws.append([line])
    help_ws["A1"].font = Font(bold=True, size=14, color=MERLOT)
    help_ws["A7"].font = Font(bold=True, color=MERLOT)
    help_ws.column_dimensions["A"].width = 110
    wb.save(out)


if __name__ == "__main__":
    build(sys.argv[1])
